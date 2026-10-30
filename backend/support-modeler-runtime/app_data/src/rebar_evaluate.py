"""N1/N2十五个建模字段评价；真值只在识别结果保存后读取。"""
from pathlib import Path
from openpyxl import load_workbook
from extractor.steel_arch import parse_angle

FIELDS = ["N1_diameter_mm", "N1_length_mm", *[f"N1_R{i}_mm" for i in range(1,5)],
          "N2_diameter_mm", "N2_length_mm", *[f"N2_r{i}_mm" for i in range(1,5)],
          "angle_A", "angle_B", "angle_C"]


def read_truth(path, support_type):
    book = load_workbook(Path(path), read_only=True, data_only=True)
    sheet = book.active
    headers = [cell.value for cell in next(sheet.iter_rows())]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row.get("support_type") == support_type:
            book.close()
            return row
    book.close()
    raise KeyError(f"ground truth中没有{support_type}")


def evaluate_rebar(result, truth):
    details = {}
    for field in FIELDS:
        item = result["evidence"].get(field)
        actual = item.get("value") if item else None
        expected = truth.get(field)
        if field.startswith("angle_"):
            expected = truth.get(field + "_deg")
            if expected is None and truth.get(field + "_raw"):
                expected = parse_angle(str(truth[field + "_raw"]))
            correct = actual is not None and expected is not None and abs(float(actual)-float(expected)) <= 0.0001
        else:
            correct = actual == expected
        details[field] = {"expected": expected, "actual": actual, "correct": correct,
                          "recognized_anywhere": any(c.get("value") == expected for group in result["candidates"].values() for c in group)}
    groups = {
        "diameter_accuracy": ["N1_diameter_mm", "N2_diameter_mm"],
        "length_accuracy": ["N1_length_mm", "N2_length_mm"],
        "radius_accuracy": [f"N1_R{i}_mm" for i in range(1,5)] + [f"N2_r{i}_mm" for i in range(1,5)],
        "angle_accuracy": ["angle_A", "angle_B", "angle_C"],
        "overall_field_accuracy": FIELDS,
        "parameter_assignment_accuracy": FIELDS,
    }
    metrics = {}
    for name, fields in groups.items():
        correct = sum(details[field]["correct"] for field in fields)
        metrics[name] = {"correct": correct, "total": len(fields), "accuracy": correct/len(fields)}
    return {"support_type": result["support_type"], "metrics": metrics, "fields": details,
            "evaluation_status": "ALL_CORRECT" if metrics["overall_field_accuracy"]["accuracy"] == 1 else "FIELD_MISMATCH"}

