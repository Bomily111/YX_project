"""运行时导出两张表；只导出已提取结果，字段评价由独立评价模块提供。"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pdf_export import save_json

STEEL_COLUMNS = ['support_type','page_numbers','has_steel_arch','steel_arch_status','R1_mm','R2_mm','beta_raw','beta_deg',
                 'steel_profile','R1_source_page','R1_source_text','R2_source_page','R2_source_text','beta_source_page',
                 'beta_source_text','steel_profile_source_page','steel_profile_source_text','evaluation_status']
PRESENCE_COLUMNS = ['support_type','page_numbers','has_steel_arch','steel_arch_pages','has_n1n2_rebar','rebar_pages','notes']

def rows_for_workbook(results, evaluation):
    statuses = {r['support_type']:r['evaluation_status'] for r in evaluation['types']}
    steel, presence = [], []
    for r in results:
        row = {k:r.get(k) for k in STEEL_COLUMNS}
        row['page_numbers'] = '、'.join(map(str,r['page_numbers']))
        row['evaluation_status'] = statuses[r['support_type']]
        for key in ('R1','R2','beta','steel_profile'):
            item = r['evidence'].get(key) or {}
            row[f'{key}_source_page'] = item.get('page')
            row[f'{key}_source_text'] = item.get('source_text')
        steel.append([row[k] for k in STEEL_COLUMNS])
        note = f"钢筋状态：{r['rebar_status']}。" + ('全部关联页未检出钢架组合安装图标题。' if not r['has_steel_arch'] else '钢架标题证据见JSON。')
        presence.append([r['support_type'],row['page_numbers'],r['has_steel_arch'],'、'.join(map(str,r['steel_arch_pages'])),
                         r['has_n1n2_rebar'],'、'.join(map(str,r['rebar_pages'])),note])
    return steel,presence

def export_supports(results, evaluation, folder):
    folder = Path(folder)
    folder.mkdir(parents=True,exist_ok=True)
    save_json(folder/'support_parameters.json',{'results':results,'evaluation':evaluation,
              'coordinate_system':'bbox: native rendered page pixels, origin top-left; page numbers one-based'})
    save_json(folder/'support_evaluation.json',evaluation)
    steel,presence = rows_for_workbook(results,evaluation)
    book = Workbook()
    book.remove(book.active)
    for name,columns,rows in [('steel_arch',STEEL_COLUMNS,steel),('component_presence',PRESENCE_COLUMNS,presence)]:
        sheet=book.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes='C2'
        sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
            cell.fill=PatternFill('solid',fgColor='334155')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=40
        for i,column in enumerate(columns,1):
            sheet.column_dimensions[get_column_letter(i)].width=40 if ('text' in column or column=='notes') else 26
        for rr in sheet.iter_rows(min_row=2):
            for cell in rr:
                cell.font=Font(name='微软雅黑',size=11)
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if cell.column_letter=='H' and name=='steel_arch':
                    cell.number_format='0.0000'
            sheet.row_dimensions[rr[0].row].height=64
    book.save(folder/'support_parameters.xlsx')
    book.close()
    lines=[]
    for r in results:
        lines.extend([r['support_type'],f"来源页：{r['page_numbers']}",f"钢架：{r['steel_arch_status']}",
                      f"R1={r['R1_mm']} mm; R2={r['R2_mm']} mm; beta={r['beta_deg']}°; profile={r['steel_profile']}",
                      f"N1/N2大样：{r['rebar_status']}，页码={r['rebar_pages']}",''])
    (folder/'support_parameters.txt').write_text('\n'.join(lines),encoding='utf-8')
