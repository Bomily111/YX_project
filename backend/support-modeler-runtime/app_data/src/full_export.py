"""全量盲提取Excel、JSON和自动摘要。"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

PRESENCE=["support_type","page_numbers","has_steel_arch","steel_arch_pages","steel_arch_status",
          "has_n1n2_rebar","rebar_pages","rebar_status"]
STEEL=["support_type","source_pages","R1_mm","R2_mm","beta_raw","beta_deg","steel_profile","overall_status"]
REBAR=["support_type","source_page","N1_diameter_mm","N1_length_mm",*[f"N1_R{i}_mm" for i in range(1,5)],
       "N2_diameter_mm","N2_length_mm",*[f"N2_r{i}_mm" for i in range(1,5)],
       "angle_A_raw","angle_A_deg","angle_B_raw","angle_B_deg","angle_C_raw","angle_C_deg",
       "geometry_consistency_status","overall_status"]
REVIEW=["support_type","component","field","predicted_value","status","reason","source_page","source_text"]


def join_pages(values):
    return "、".join(map(str,values or []))


def value(component,field):
    return component.get("fields",{}).get(field,{}).get("value")


def raw(component,field):
    return component.get("fields",{}).get(field,{}).get("raw_text")


def workbook_rows(payload):
    presence=[];steel=[];rebar=[]
    for item in payload["supports"]:
        s=item["steel_arch"];r=item["rebar"]
        presence.append([item["support_type"],join_pages(item["page_numbers"]),s.get("has_component",False),
                         join_pages(s.get("source_pages",[])),s.get("status"),
                         bool(r and r.get("has_component")),join_pages(r.get("source_pages",[]) if r else []),
                         r.get("status") if r else "NOT_APPLICABLE"])
        if s.get("has_component"):
            steel.append([item["support_type"],join_pages(s.get("source_pages")),value(s,"R1_mm"),value(s,"R2_mm"),
                          raw(s,"beta_deg"),value(s,"beta_deg"),value(s,"steel_profile"),s["status"]])
        if r and r.get("has_component"):
            rebar.append([item["support_type"],join_pages(r.get("source_pages")),value(r,"N1_diameter_mm"),
                          value(r,"N1_length_mm"),*[value(r,f"N1_R{i}_mm") for i in range(1,5)],
                          value(r,"N2_diameter_mm"),value(r,"N2_length_mm"),
                          *[value(r,f"N2_r{i}_mm") for i in range(1,5)],
                          raw(r,"angle_A"),value(r,"angle_A"),raw(r,"angle_B"),value(r,"angle_B"),
                          raw(r,"angle_C"),value(r,"angle_C"),(r.get("geometry_consistency") or {}).get("status"),r["status"]])
    review=[[row.get(column) for column in REVIEW] for row in payload["review"]]
    return presence,steel,rebar,review


def style_sheet(sheet,columns):
    sheet.freeze_panes="C2"
    sheet.auto_filter.ref=sheet.dimensions
    sheet.sheet_view.showGridLines=False
    for cell in sheet[1]:
        cell.font=Font(name="微软雅黑",bold=True,color="FFFFFF",size=11)
        cell.fill=PatternFill("solid",fgColor="334155")
        cell.alignment=Alignment(wrap_text=True,vertical="center")
    sheet.row_dimensions[1].height=42
    for index,column in enumerate(columns,1):
        sheet.column_dimensions[get_column_letter(index)].width=32 if column in ("reason","source_text") else 22
        for cell in sheet.iter_cols(min_col=index,max_col=index,min_row=2):
            for value_cell in cell:
                value_cell.font=Font(name="微软雅黑",size=10)
                value_cell.alignment=Alignment(wrap_text=True,vertical="top")
        if column.endswith("_deg"):
            for cell in sheet.iter_cols(min_col=index,max_col=index,min_row=2):
                for value_cell in cell: value_cell.number_format="0.0000"


def summary_markdown(payload):
    supports=payload["supports"]
    found=[s for s in supports if s["page_numbers"]]
    steel_yes=[s["support_type"] for s in supports if s["steel_arch"].get("has_component")]
    steel_no=[s["support_type"] for s in supports if not s["steel_arch"].get("has_component")]
    rebar_yes=[s["support_type"] for s in supports if s["rebar"] and s["rebar"].get("has_component")]
    rebar_no=[s["support_type"] for s in supports if not s["rebar"] or not s["rebar"].get("has_component")]
    schema=sum(1 for s in supports for c in (s["steel_arch"],s["rebar"] or {}) if c.get("status")=="SCHEMA_MISMATCH")
    steel_ok=sum(s["steel_arch"].get("status")=="EXTRACTED" for s in supports)
    rebar_ok=sum(bool(s["rebar"]) and s["rebar"].get("status")=="EXTRACTED" for s in supports)
    lines=["# 完整支护图纸批量盲提取摘要","",f"- PDF总页数：{payload['metadata']['page_count']}",
           f"- 找到目标类型：{len(found)}/{len(supports)}",f"- 成功钢架参数组：{steel_ok}",
           f"- 成功钢筋参数组：{rebar_ok}",f"- SCHEMA_MISMATCH：{schema}",
           f"- 需要人工复核字段/项目：{len(payload['review'])}",
           f"- 基线回归：{'全部通过' if payload['regression']['all_passed'] else '失败'}","",
           "## 页面范围","", "| support_type | 页码 |","|---|---|"]
    lines += [f"| {s['support_type']} | {join_pages(s['page_numbers']) or '未定位'} |" for s in supports]
    lines += ["","## 构件存在性","",f"- 存在钢架：{'、'.join(steel_yes) or '无'}",
              f"- 不存在钢架：{'、'.join(steel_no) or '无'}",f"- 存在N1/N2大样：{'、'.join(rebar_yes) or '无'}",
              f"- 不存在N1/N2大样：{'、'.join(rebar_no) or '无'}","","## Review摘要",""]
    if payload["review"]:
        for row in payload["review"]:
            lines.append(f"- {row['support_type']} / {row['component']} / {row['field']}：{row['reason']}")
    else:
        lines.append("- 无")
    return "\n".join(lines)+"\n"


def export_all(payload,folder):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    (folder/"support_parameters_all.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    rows=workbook_rows(payload)
    book=Workbook();book.remove(book.active)
    for name,columns,data in zip(("component_presence","steel_arch_parameters","rebar_parameters","review_items"),
                                 (PRESENCE,STEEL,REBAR,REVIEW),rows):
        sheet=book.create_sheet(name);sheet.append(columns)
        for row in data:sheet.append(row)
        style_sheet(sheet,columns)
    book.save(folder/"support_parameters_all.xlsx");book.close()
    (folder/"full_run_summary.md").write_text(summary_markdown(payload),encoding="utf-8")
