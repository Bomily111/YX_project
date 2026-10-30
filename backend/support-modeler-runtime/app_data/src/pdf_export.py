"""Python运行时的Excel/TXT/JSON输出与独立真值评价。"""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def save_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')

def flatten(result):
    row = {k: result.get(k) for k in ('source_pdf', 'support_type', 'page_numbers', 'R1_mm', 'R2_mm', 'beta_raw', 'beta_deg', 'steel_profile')}
    row['page_numbers'] = '、'.join(map(str, result['page_numbers']))
    for key in ('R1', 'R2', 'beta', 'steel_profile'):
        item = result['evidence'].get(key) or {}
        for field in ('source_text', 'page', 'bbox'):
            row[f'{key}_{field}'] = item.get(field)
    return row

def export_result(result, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    save_json(folder / 'steel_arch_parameters.json', result)
    row = flatten(result)
    wb = Workbook()
    ws = wb.active
    ws.title = '钢架参数'
    ws.append(list(row))
    ws.append([json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in row.values()])
    for cell in ws[1]:
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='334155')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for cell in ws[2]:
        cell.font = Font(name='微软雅黑', size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for index, key in enumerate(row, 1):
        ws.column_dimensions[get_column_letter(index)].width = 36 if 'bbox' in key or 'text' in key else 23
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 65
    ws['G2'].number_format = '0.0000'
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    notes = wb.create_sheet('说明')
    for values in [('项目', '说明'), ('来源', result['source_pdf']),
                   ('页码', 'PDF物理页码，从1开始；不等同图纸印刷页码'),
                   ('坐标', f"bbox为{result['dpi']} DPI原生高清页面像素矩形[x0,y0,x1,y1]，左上角为原点"),
                   ('单位', '半径mm，角度degree；JSON另存PDF点坐标'),
                   ('缺值', '空单元格表示未确认；不使用真值补值'),
                   ('beta', '由下部空间位置和左右对称一致性确定；候选见JSON')]:
        notes.append(values)
    notes.column_dimensions['A'].width = 18
    notes.column_dimensions['B'].width = 95
    for rr in notes:
        for c in rr:
            c.font = Font(name='微软雅黑', size=11)
            c.alignment = Alignment(wrap_text=True, vertical='top')
        notes.row_dimensions[rr[0].row].height = 32
    wb.save(folder / 'steel_arch_parameters.xlsx')
    wb.close()
    lines = [f"支护类型：{result['support_type']}", f"来源PDF：{result['source_pdf']}",
             f"来源页：{row['page_numbers']}", '', f"R1：{result['R1_mm']} mm", f"R2：{result['R2_mm']} mm", '',
             'beta：', str(result['beta_raw']), f"{result['beta_deg']}°", '', '工字钢型号：', str(result['steel_profile']), '',
             f"坐标：{result['dpi']} DPI页面像素，左上角原点；页码从1开始"]
    for key, item in result['evidence'].items():
        lines.append(f"{key}证据：" + (f"第{item['page']}页，{item['source_text']}，bbox={item['bbox']}" if item else '未确认'))
    (folder / 'steel_arch_parameters.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')

def evaluate_xlsx(result, truth_path):
    """只能在识别输出完成后调用；此模块不向识别器提供任何真值。"""
    wb = load_workbook(Path(truth_path), read_only=True, data_only=True)
    try:
        for sheet in wb:
            values = list(sheet.values)
            if not values:
                continue
            headers = [str(v).strip() for v in values[0]]
            for row in values[1:]:
                if result['support_type'] not in row:
                    continue
                raw = dict(zip(headers, row))
                expected = {
                    'R1_mm': raw.get('R1_mm', raw.get('R1')),
                    'R2_mm': raw.get('R2_mm', raw.get('R2')),
                    'beta_deg': raw.get('beta_deg', raw.get('beta')),
                    'steel_profile': raw.get('steel_profile', raw.get('工字钢', raw.get('工字钢型号')))}
                report = []
                for key, exp in expected.items():
                    actual = result[key]
                    if exp is None:
                        status = '未标注'
                    elif actual is None:
                        status = '未识别'
                    elif key == 'beta_deg':
                        status = '正确' if abs(float(exp) - actual) <= .00005 else '错误'
                    else:
                        status = '正确' if exp == actual else '错误'
                    report.append({'field': key, 'expected': exp, 'actual': actual, 'status': status})
                return {'fields': report, 'all_correct': all(r['status'] == '正确' for r in report)}
        return {'fields': [], 'all_correct': False, 'error': '未找到目标支护类型真值'}
    finally:
        wb.close()
