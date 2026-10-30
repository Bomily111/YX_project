"""盲测预测输出：不读取ground truth，不计算准确率。"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def prediction_payload(result):
    return {
        "support_type": result["support_type"],
        "source_page": result["source_page"],
        "has_n1n2_rebar": result["has_n1n2_rebar"],
        "rebar_status": result["rebar_status"],
        "N1": result["N1"],
        "N2": result["N2"],
        "angles": result["angles"],
        "geometry_consistency": result["geometry_consistency"],
        "missing_fields": result["missing_fields"],
        "field_sources": {
            field: {
                "value": item["value"],
                "source_text": item["source_text"],
                "source_method": item["source_method"],
                "confidence": item["confidence"],
                "page": item["page"],
                "bbox": item["bbox"],
            }
            for field, item in result["evidence"].items()
        },
    }


def export_prediction(result, folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    payload = prediction_payload(result)
    (folder/"prediction.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    columns = ["kind", "owner", "raw_name", "normalized_name", "value", "source_text",
               "source_method", "page", "bbox", "confidence", "rotation", "stage"]
    rows = []
    for kind, candidates in result["candidates"].items():
        for item in candidates:
            row = dict(item)
            row["kind"] = kind
            row["bbox"] = json.dumps(item.get("bbox"), ensure_ascii=False)
            values = []
            for column in columns:
                value = row.get(column)
                # OCR原文可能以=开头；作为文本保存，禁止Excel按公式执行。
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    value = "'" + value
                values.append(value)
            rows.append(values)
    book = Workbook()
    sheet = book.active
    sheet.title = "candidates"
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="334155")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 30 if column in ("source_text","bbox") else 18
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(folder/"candidate_parameters.xlsx")
    book.close()
    return payload
